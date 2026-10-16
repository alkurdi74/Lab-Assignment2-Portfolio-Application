from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml.ns import qn
from PIL import Image
from pathlib import Path
base = Path(__file__).resolve().parents[1]
doc = Document()
section = doc.sections[0]
section.top_margin = section.bottom_margin = Inches(.65)
section.left_margin = section.right_margin = Inches(.7)
for name in ['Normal', 'Title', 'Heading 1', 'Heading 2']:
    doc.styles[name].font.name = 'Calibri'
    doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)
doc.styles['Normal'].font.size = Pt(11)
doc.add_heading('Portfolio Backend Assignment 2', 0)
doc.add_paragraph('COMP229 Web Application Development | Mohamad Alkurdi')
doc.add_paragraph('The portfolio backend implements Node.js and Express REST APIs with Mongoose, the Portfolio MongoDB database, and JWT authentication. All 24 required CRUD endpoints are implemented. The existing React contact form saves submissions to MongoDB.')
doc.add_heading('Verification results', 1)
doc.add_paragraph('79 HTTP requests passed against a real, isolated MongoDB test database. The checks cover all CRUD operations, input validation, missing records, duplicate users, hashed passwords, protected routes, password changes, and token revocation after sign-out. Frontend production build and ESLint both passed.')
doc.add_heading('Evidence method', 1)
doc.add_paragraph('The welcome image is a browser screenshot of the running API. The database image displays an actual Mongoose connection and collection listing. API images are browser-rendered views of captured integration-test requests and responses, not Postman screenshots. Tokens and passwords are redacted. Complete raw evidence is included in evidence/api-results.json and a runnable Postman collection is supplied.')
doc.add_heading('Running the project', 1)
doc.add_paragraph('Install dependencies in the root and client folders, configure .env from .env.example, start MongoDB, and run npm run dev. The frontend uses port 5173 and the backend uses port 3000. See README.md for commands and endpoint access rules.')
doc.add_heading('Submission checklist', 1)
doc.add_paragraph('Include the project ZIP, this Word report, and your GitHub repository URL. A GitHub repository has not been published from this workspace; add its real URL when submitting. The database and console figures display captured MongoDB results and server process output.')
images = [('welcome','Backend welcome response'),('database','MongoDB database and collections'),('console','Backend startup console')]
for resource in ['contacts','projects','qualifications','users']:
    images.extend([(resource+'-create',resource.title()+' create request and response'),(resource+'-summary',resource.title()+' CRUD and validation checks')])
images.append(('authentication','Authentication and sign out checks'))
for name, title in images:
    doc.add_page_break()
    doc.add_heading(title, 1)
    p = base / 'evidence' / (name+'.png')
    w,h = Image.open(p).size
    width = min(7.0, 8.6*w/h)
    doc.add_picture(str(p), width=Inches(width))
    doc.add_paragraph('Captured from the local running backend and test results on 13 October 2026.')
for style in doc.styles:
    for border in list(style.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
for border in list(doc.element.iter(qn('w:pBdr'))):
    border.getparent().remove(border)
out = base.parent / 'Assignment_2_Screenshot_Report.docx'
doc.save(out)
print(out)
